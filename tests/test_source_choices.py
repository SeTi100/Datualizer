"""Choosing a sheet, table or clock date when a source needs it (Etappe 2b, follow-up to step 3).

A lab journal has one sheet per measuring day and only times of day; a SQLite file without a
profile can hold several tables. The reader reports the choices machine-readably, a profile
copy stores the user's pick, and the campaign replays it exactly.
Multi-sheet workbooks and multi-table databases are built in tmp_path; fixtures stay untouched.
"""

from __future__ import annotations

from datetime import date, datetime
import json
from pathlib import Path
import sqlite3

import openpyxl
import pytest

from datualizer_core import (
    ReaderError,
    date_in_text,
    detect_profile,
    load_source,
    needs_clock_date,
    with_choices,
)
from datualizer_core.campaign import Campaign, CampaignSession
from datualizer_core.ingestion.readers import read_excel, read_sqlite

SA = Path(__file__).parent / "fixtures" / "sa_12_10"
JOURNAL = SA / "labjournal_12_10_extract.xlsx"
DB = SA / "instrument_reader_12_10_extract.db"
SHEETS = ["12.10.2026", "13.10.2026"]


def make_multi_sheet_journal(path: Path) -> Path:
    """Journal workbook with the fixture's sheet twice, as on two measuring days."""
    src = openpyxl.load_workbook(JOURNAL, read_only=True, data_only=True)
    rows = [list(r) for r in src.worksheets[0].iter_rows(values_only=True)]
    src.close()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in SHEETS:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


def make_multi_table_db(path: Path) -> Path:
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE a (t TEXT, x REAL)")
    con.execute("CREATE TABLE b (t TEXT, y REAL)")
    con.executemany("INSERT INTO a VALUES (?, ?)", [("2026-10-12 10:00:00", 1.0), ("2026-10-12 10:00:01", 2.0)])
    con.executemany("INSERT INTO b VALUES (?, ?)", [("2026-10-12 11:00:00", 5.0)] * 3)
    con.commit()
    con.close()
    return path


@pytest.fixture()
def journal(tmp_path: Path) -> Path:
    return make_multi_sheet_journal(tmp_path / "SA_Labjournal.xlsx")


class TestChoicesInReaderError:
    def test_several_sheets_list_the_choice(self, journal: Path) -> None:
        with pytest.raises(ReaderError, match="several sheets") as info:
            read_excel(journal)
        assert info.value.choice == "excel.sheet"
        assert info.value.options == SHEETS

    def test_several_tables_list_the_choice(self, tmp_path: Path) -> None:
        with pytest.raises(ReaderError) as info:
            read_sqlite(make_multi_table_db(tmp_path / "x.db"))
        assert info.value.choice == "sqlite.table"
        assert info.value.options == ["a", "b"]
        with pytest.raises(ReaderError) as info:
            read_sqlite(DB)
        assert info.value.choice == "sqlite.table"
        assert "sqlite_sequence" not in info.value.options

    def test_load_source_passes_the_choice_on(self, journal: Path) -> None:
        assert detect_profile(journal).name == "labjournal_xlsx"
        with pytest.raises(ReaderError) as info:
            load_source(journal)
        assert info.value.options == SHEETS


class TestDateInText:
    @pytest.mark.parametrize(
        "text, expected",
        [
            ("12.10.2026", date(2026, 10, 12)),
            ("Messung 5.10.2026 kalt", date(2026, 10, 5)),
            ("2026-10-13", date(2026, 10, 13)),
            ("SA_12_10_2026.xlsx", date(2026, 10, 12)),
            ("Tabelle1", None),
            ("32.13.2026", None),
            ("", None),
        ],
    )
    def test_suggestion(self, text: str, expected: date | None) -> None:
        assert date_in_text(text) == expected


class TestWithChoices:
    def test_sheet_and_clock_date_place_the_journal(self, journal: Path) -> None:
        detected = detect_profile(journal)
        prof = with_choices(journal, detected, sheet="13.10.2026", clock_date=date(2026, 10, 13))
        assert detected.excel.sheet is None and detected.ingestion.time.clock_date is None  # copy
        ds = load_source(journal, prof)
        assert ds.source_meta["sheet"] == "13.10.2026"
        assert ds.time_abs.dt.replace_time_zone(None)[1] == datetime(2026, 10, 13, 10, 53)
        assert not needs_clock_date(ds)

    def test_without_date_the_journal_needs_one(self, journal: Path) -> None:
        ds = load_source(journal, with_choices(journal, detect_profile(journal), sheet="12.10.2026"))
        assert needs_clock_date(ds)
        assert ds.time_origin is None

    def test_sqlite_without_profile(self, tmp_path: Path) -> None:
        db = make_multi_table_db(tmp_path / "x.db")
        assert detect_profile(db) is None
        prof = with_choices(db, None, table="b")
        ds = load_source(db, prof)
        assert len(ds) == 3
        assert ds.source_meta["profile"] == "manual" and ds.source_meta["reader"] == "sqlite"
        assert not needs_clock_date(ds)

    def test_table_replaces_a_query(self) -> None:
        detected = detect_profile(DB)
        prof = with_choices(DB, detected, table="runs")
        assert prof.sqlite.table == "runs" and prof.sqlite.query is None
        assert detected.sqlite.query is not None


class TestCampaignReplay:
    def test_campaign_json_reloads_the_chosen_sheet_and_date(self, journal: Path, tmp_path: Path) -> None:
        session = CampaignSession(Campaign())
        prof = with_choices(journal, detect_profile(journal), sheet="13.10.2026", clock_date=date(2026, 10, 13))
        ds = load_source(journal, prof)
        src = session.add_source(journal, prof, dataset=ds)
        out = tmp_path / "campaign.json"
        session.save(out)

        stored = json.loads(out.read_text(encoding="utf-8"))["sources"][0]
        assert stored["profile"]["excel"]["sheet"] == "13.10.2026"
        assert stored["spec"]["time"]["clock_date"] == "2026-10-13"

        reopened = CampaignSession.open(out)
        assert reopened.warnings == []
        again = reopened.datasets[src.id]
        assert again.df.equals(ds.df)
        assert again.time_origin == ds.time_origin
        assert reopened.span(src.id) == session.span(src.id)
