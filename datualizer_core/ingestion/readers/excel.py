"""Excel workbooks: `.xlsx` via openpyxl, legacy `.xls` (BIFF) via xlrd.

One sheet becomes the data table. Small key/value sheets (e.g. device settings and calibration
of an analyzer export) become source metadata instead of being lost.
"""

from __future__ import annotations

from pathlib import Path
import re
from typing import Any, Sequence

from pydantic import BaseModel, ConfigDict, Field

from datualizer_core.ingestion.readers.base import RawTable, ReaderError, cell_to_text, string_frame, unique_headers

HEADER_SCAN_ROWS = 50


class ExcelOptions(BaseModel):
    """Which part of a workbook to read."""

    model_config = ConfigDict(extra="forbid")

    sheet: str | int | None = Field(
        None, description="Sheet name or 0-based index. None = the only sheet, else an error lists them."
    )
    header_row: int | None = Field(
        None, ge=1, description="1-based header row. None = the row with the most text cells among the first 50."
    )
    meta_sheets: list[str] | None = Field(
        None,
        description=(
            "Sheets read as key/value metadata. None = every other sheet with at most 3 used columns."
        ),
    )


def _column_letter(idx: int) -> str:
    letters = ""
    idx += 1
    while idx:
        idx, rem = divmod(idx - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _rows_xlsx(path: Path) -> tuple[list[str], dict[str, list[list[Any]]]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    finally:
        wb.close()
    return list(sheets), sheets


def _rows_xls(path: Path) -> tuple[list[str], dict[str, list[list[Any]]]]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    sheets: dict[str, list[list[Any]]] = {}
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append(xlrd.error_text_from_code.get(cell.value, "#ERROR"))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(bool(cell.value))
                else:
                    row.append(cell.value)
            rows.append(row)
        sheets[sh.name] = rows
    return book.sheet_names(), sheets


def _used_width(rows: list[list[Any]]) -> int:
    width = 0
    for r in rows:
        for i in range(len(r) - 1, -1, -1):
            if r[i] is not None and str(r[i]).strip() != "":
                width = max(width, i + 1)
                break
    return width


def _detect_header(rows: list[list[Any]]) -> int:
    """0-based index of the row with the most text cells (followed by a non-empty row)."""
    best, best_n = 0, -1
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        n = sum(isinstance(v, str) and v.strip() != "" for v in row)
        has_next = i + 1 < len(rows) and any(v is not None and str(v).strip() for v in rows[i + 1])
        if n > best_n and has_next:
            best, best_n = i, n
    return best


def _meta_from_sheet(rows: list[list[Any]]) -> dict[str, str]:
    meta: dict[str, str] = {}
    for r in rows:
        if len(r) >= 2 and r[0] is not None and r[1] is not None:
            meta[str(cell_to_text(r[0])).strip()] = str(cell_to_text(r[1])).strip()
    return meta


def _workbook_rows(path: Path) -> tuple[list[str], dict[str, list[list[Any]]]]:
    return _rows_xls(path) if path.suffix.lower() == ".xls" else _rows_xlsx(path)


def sheet_names(path: str | Path) -> list[str]:
    """Sheet names of a workbook, without reading the cells."""
    path = Path(path)
    if path.suffix.lower() == ".xls":
        import xlrd

        return xlrd.open_workbook(str(path), on_demand=True).sheet_names()
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _pick_sheet(path: Path, names: list[str], sheet: str | int | None) -> str:
    """Name of the requested sheet; a ReaderError with the sheet names as choice otherwise."""
    if sheet is None:
        if len(names) != 1:
            raise ReaderError(
                f"{path.name} has several sheets, choose one of: {names}", choice="excel.sheet", options=names
            )
        return names[0]
    if isinstance(sheet, int):
        if not 0 <= sheet < len(names):
            raise ReaderError(
                f"{path.name} has no sheet #{sheet}; sheets: {names}", choice="excel.sheet", options=names
            )
        return names[sheet]
    if sheet in names:
        return sheet
    raise ReaderError(f"{path.name} has no sheet {sheet!r}; sheets: {names}", choice="excel.sheet", options=names)


_CELL = re.compile(r"^\s*([A-Za-z]{1,3})\s*(\d+)\s*$")


def parse_cell(ref: str) -> tuple[int, int]:
    """0-based (row, column) of an A1 reference such as 'C3'; ValueError otherwise."""
    m = _CELL.match(ref)
    if m is None or int(m.group(2)) < 1:
        raise ValueError(f"Not a cell reference: {ref!r} (expected e.g. 'C3').")
    col = 0
    for ch in m.group(1).upper():
        col = col * 26 + ord(ch) - 64
    return int(m.group(2)) - 1, col - 1


def read_cells(path: str | Path, sheet: str | int | None, cells: Sequence[str]) -> dict[str, Any]:
    """Raw values of single cells (A1 references) of one sheet; empty cells are None.

    Values come as Excel stores them (number, text, date-time, time), not as text, so the caller
    decides how to read them.
    """
    path = Path(path)
    names, sheets = _workbook_rows(path)
    rows = sheets[_pick_sheet(path, names, sheet)]
    out: dict[str, Any] = {}
    for ref in cells:
        r, c = parse_cell(ref)
        out[ref] = rows[r][c] if r < len(rows) and c < len(rows[r]) else None
    return out


def read_sheet(path: str | Path, sheet: str | int | None) -> tuple[str, list[str], list[list[Any]], int]:
    """All rows of one sheet with typed values (also those above the header).

    Returns the sheet name, all sheet names, the rows (row 1 first) and the 1-based row the
    reader would take as header. Only this sheet is read (.xlsx), so it is quick for a large journal.
    """
    path = Path(path)
    if path.suffix.lower() == ".xls":
        names, sheets = _rows_xls(path)
        name = _pick_sheet(path, names, sheet)
        rows = sheets[name]
    else:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            names = list(wb.sheetnames)
            name = _pick_sheet(path, names, sheet)
            rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
        finally:
            wb.close()
    while rows and all(v is None or str(v).strip() == "" for v in rows[-1]):
        rows.pop()
    return name, names, rows, _detect_header(rows) + 1 if rows else 1


def read_excel(path: str | Path, options: ExcelOptions | None = None) -> RawTable:
    """Read one sheet of a workbook as an all-String table; small sheets become metadata."""
    path = Path(path)
    opts = options or ExcelOptions()
    names, sheets = _workbook_rows(path)
    sheet = _pick_sheet(path, names, opts.sheet)

    rows = sheets[sheet]
    width = _used_width(rows)
    if width == 0:
        raise ReaderError(f"Sheet {sheet!r} of {path.name} is empty.")
    header_idx = opts.header_row - 1 if opts.header_row is not None else _detect_header(rows)
    if header_idx >= len(rows):
        raise ReaderError(f"Sheet {sheet!r} has no row {header_idx + 1}.")
    header = (rows[header_idx] + [None] * width)[:width]
    headers = unique_headers(header, [f"col_{_column_letter(i)}" for i in range(width)])
    data = [r[:width] for r in rows[header_idx + 1:]]
    while data and all(v is None or str(v).strip() == "" for v in data[-1]):
        data.pop()

    meta_names = opts.meta_sheets
    if meta_names is None:
        meta_names = [n for n in names if n != sheet and 0 < _used_width(sheets[n]) <= 3]
    missing = [n for n in meta_names if n not in sheets]
    if missing:
        raise ReaderError(f"{path.name} has no metadata sheet(s) {missing}; sheets: {names}")

    meta = {
        "sheet": sheet,
        "sheets": names,
        "header_row": header_idx + 1,
        "meta_sheets": {n: _meta_from_sheet(sheets[n]) for n in meta_names},
    }
    return RawTable(string_frame(headers, data), ".", meta)
